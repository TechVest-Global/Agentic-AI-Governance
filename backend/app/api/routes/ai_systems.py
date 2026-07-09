from typing import Annotated
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, Query, UploadFile, status
from sqlmodel import Session

from app.core.exceptions import ApplicationError
from app.db.session import get_session
from app.schemas.governance import (
    AISystemCapabilityCreate,
    AISystemCapabilityRead,
    AISystemCreate,
    AISystemRead,
    AISystemRegistrationCreate,
    AISystemRegistrationRead,
    AISystemUpdate,
    ApplicationContextProfileCreate,
    ApplicationContextProfileRead,
    RetrievalContextDocumentCreate,
    RetrievalContextDocumentRead,
)
from app.services import ai_systems as service
from app.services import registration as registration_service

router = APIRouter(prefix="/ai-systems")
SessionDependency = Annotated[Session, Depends(get_session)]


@router.post("", response_model=AISystemRead, status_code=status.HTTP_201_CREATED)
def create_ai_system(
    payload: AISystemCreate,
    session: SessionDependency,
) -> AISystemRead:
    return service.create_ai_system(session, payload)


@router.post(
    "/register",
    response_model=AISystemRegistrationRead,
    status_code=status.HTTP_201_CREATED,
)
def register_ai_system(
    payload: AISystemRegistrationCreate,
    session: SessionDependency,
) -> AISystemRegistrationRead:
    """Register a complete AI system (nested facts) and derive a preliminary risk tier."""
    return registration_service.register_ai_system(session, payload)


@router.get("", response_model=list[AISystemRead])
def list_ai_systems(
    session: SessionDependency,
    offset: Annotated[int, Query(ge=0)] = 0,
    limit: Annotated[int, Query(ge=1, le=100)] = 50,
) -> list[AISystemRead]:
    return service.list_ai_systems(session, offset=offset, limit=limit)


@router.get("/{system_id}", response_model=AISystemRead)
def get_ai_system(system_id: UUID, session: SessionDependency) -> AISystemRead:
    return service.get_ai_system(session, system_id)


@router.get("/{system_id}/registration", response_model=AISystemRegistrationRead)
def get_ai_system_registration(
    system_id: UUID, session: SessionDependency
) -> AISystemRegistrationRead:
    """Return the complete nested registration record for an AI system."""
    return registration_service.get_registration_detail(session, system_id)


@router.patch("/{system_id}", response_model=AISystemRead)
def update_ai_system(
    system_id: UUID,
    payload: AISystemUpdate,
    session: SessionDependency,
) -> AISystemRead:
    return service.update_ai_system(session, system_id, payload)


@router.delete("/{system_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_ai_system(system_id: UUID, session: SessionDependency) -> None:
    service.archive_ai_system(session, system_id)


@router.post(
    "/{system_id}/capabilities",
    response_model=AISystemCapabilityRead,
    status_code=status.HTTP_201_CREATED,
)
def create_capability(
    system_id: UUID,
    payload: AISystemCapabilityCreate,
    session: SessionDependency,
) -> AISystemCapabilityRead:
    return service.create_capability(session, system_id, payload)


@router.post("/{system_id}/capabilities/import-from-catalog")
def import_capabilities_from_catalog(
    system_id: UUID,
    session: SessionDependency,
) -> dict:
    """Create one capability per endpoint in the target's gateway catalog.

    Fetches ``GET {target_endpoint_ref}/catalog`` using the configured target API
    key and imports every endpoint as a capability (idempotent). Lets an auditor
    register all of a multi-endpoint system's functions in one click instead of
    entering them by hand.
    """
    from dataclasses import asdict as _asdict  # noqa: F401  (kept for parity)

    from app.core.config import get_settings
    from app.services.capability_import import import_capabilities_from_catalog as _import

    result = _import(session, system_id, api_key=get_settings().target_api_key)
    # Serialize capability rows for the response.
    result["capabilities"] = [
        {
            "id": str(c.id),
            "name": c.name,
            "endpoint_ref": c.endpoint_ref,
            "http_method": c.http_method,
        }
        for c in result["capabilities"]
    ]
    return result


@router.get(
    "/{system_id}/capabilities",
    response_model=list[AISystemCapabilityRead],
)
def list_capabilities(
    system_id: UUID,
    session: SessionDependency,
    offset: Annotated[int, Query(ge=0)] = 0,
    limit: Annotated[int, Query(ge=1, le=100)] = 50,
    enabled: bool | None = None,
) -> list[AISystemCapabilityRead]:
    return service.list_capabilities(
        session,
        system_id,
        offset=offset,
        limit=limit,
        enabled=enabled,
    )


@router.get(
    "/{system_id}/capabilities/{capability_id}",
    response_model=AISystemCapabilityRead,
)
def get_capability(
    system_id: UUID,
    capability_id: UUID,
    session: SessionDependency,
) -> AISystemCapabilityRead:
    return service.get_capability(session, system_id, capability_id)


@router.get(
    "/{system_id}/context-profile",
    response_model=ApplicationContextProfileRead,
)
def get_context_profile(
    system_id: UUID,
    session: SessionDependency,
) -> ApplicationContextProfileRead:
    return service.get_context_profile(session, system_id)


@router.put(
    "/{system_id}/context-profile",
    response_model=ApplicationContextProfileRead,
)
def upsert_context_profile(
    system_id: UUID,
    payload: ApplicationContextProfileCreate,
    session: SessionDependency,
) -> ApplicationContextProfileRead:
    return service.upsert_context_profile(session, system_id, payload)


@router.post(
    "/{system_id}/retrieval-context",
    response_model=RetrievalContextDocumentRead,
    status_code=status.HTTP_201_CREATED,
)
def create_retrieval_context_document(
    system_id: UUID,
    payload: RetrievalContextDocumentCreate,
    session: SessionDependency,
) -> RetrievalContextDocumentRead:
    return service.create_retrieval_context_document(session, system_id, payload)


@router.post(
    "/{system_id}/retrieval-context/upload",
    response_model=RetrievalContextDocumentRead,
    status_code=status.HTTP_201_CREATED,
)
async def upload_retrieval_context_document(
    system_id: UUID,
    session: SessionDependency,
    file: Annotated[UploadFile | None, File()] = None,
    text: Annotated[str | None, Form()] = None,
    title: Annotated[str | None, Form()] = None,
    tags: Annotated[str | None, Form()] = None,
) -> RetrievalContextDocumentRead:
    """Ingest an uploaded file (or pasted text) as a retrieval-context document.

    Lets users add context by upload instead of hand-typing JSON. The stored
    document is what RAG groundedness evaluation (RAGAS) reads at run time.
    Accepts a multipart file (.pdf / .docx / plain text formats) or a raw
    ``text`` field; ``tags`` is a comma-separated string.
    """
    filename = None
    if file is not None:
        raw = await file.read()
        filename = file.filename
        content = _extract_document_text(raw, filename)
    else:
        content = text or ""
    tag_list = [t.strip() for t in (tags or "").split(",") if t.strip()]
    return service.ingest_context_document(
        session,
        system_id,
        title=title,
        content=content,
        source_filename=filename,
        tags=tag_list,
    )


def _extract_document_text(raw: bytes, filename: str | None) -> str:
    """Extract plain text from an uploaded document by its extension.

    PDFs and DOCX were previously decoded as UTF-8, storing binary mojibake as
    "context" that then poisoned RAG groundedness evaluation. Anything without
    a known binary extension is treated as plain text.
    """
    import io

    suffix = (filename or "").rsplit(".", 1)[-1].lower() if filename and "." in filename else ""

    if suffix == "pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            pages = [page.extract_text() or "" for page in reader.pages]
            content = "\n\n".join(part.strip() for part in pages if part.strip())
        except Exception as exc:
            raise ApplicationError(
                status_code=422,
                code="VALIDATION_ERROR",
                message=f"Could not read the PDF file: {exc}",
            ) from exc
        if not content.strip():
            raise ApplicationError(
                status_code=422,
                code="VALIDATION_ERROR",
                message=(
                    "The PDF contains no extractable text (it may be scanned images). "
                    "Paste the text instead."
                ),
            )
        return content

    if suffix in {"docx", "doc"}:
        try:
            from docx import Document

            document = Document(io.BytesIO(raw))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            content = "\n".join(parts)
        except Exception as exc:
            raise ApplicationError(
                status_code=422,
                code="VALIDATION_ERROR",
                message=(
                    f"Could not read the Word document: {exc}. Legacy .doc files are "
                    "not supported — save as .docx and retry."
                ),
            ) from exc
        if not content.strip():
            raise ApplicationError(
                status_code=422,
                code="VALIDATION_ERROR",
                message="The Word document contains no extractable text.",
            )
        return content

    return raw.decode("utf-8", errors="replace")


@router.get(
    "/{system_id}/retrieval-context",
    response_model=list[RetrievalContextDocumentRead],
)
def list_retrieval_context_documents(
    system_id: UUID,
    session: SessionDependency,
    offset: Annotated[int, Query(ge=0)] = 0,
    limit: Annotated[int, Query(ge=1, le=100)] = 50,
) -> list[RetrievalContextDocumentRead]:
    return service.list_retrieval_context_documents(
        session, system_id, offset=offset, limit=limit
    )


@router.delete(
    "/{system_id}/retrieval-context/{document_id}",
    status_code=status.HTTP_204_NO_CONTENT,
)
def delete_retrieval_context_document(
    system_id: UUID,
    document_id: UUID,
    session: SessionDependency,
) -> None:
    service.delete_retrieval_context_document(session, system_id, document_id)
