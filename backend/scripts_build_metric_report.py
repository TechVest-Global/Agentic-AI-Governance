"""Build a plain-English Excel report of how each metric tested the chatbot and
how its score was produced. 3 tabs: Overview, All Metrics, Probes & Responses.

Usage:
  python scripts_build_metric_report.py <result.json> <out.xlsx> [run_id] [system]
"""
import json
import sys
from collections import Counter

from app.configs.config_loader import load_metric_configs_from_dir
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RESULT_PATH, OUT_PATH = sys.argv[1], sys.argv[2]
RUN_ID = sys.argv[3] if len(sys.argv) > 3 else "(unknown)"
SYSTEM = sys.argv[4] if len(sys.argv) > 4 else "(unknown)"
SELECTED_FW = ["eu_ai_act", "iso_42001", "nist_ai_rmf", "owasp_llm_top_10",
               "owasp_agentic_ai", "mitre_atlas"]

# ---------- style ----------
NAVY, TEAL, TEAL_SOFT, WHITE = "16202A", "0E7C6B", "E2F0EC", "FFFFFF"
GREEN, RED, AMBER, HAIR, SLATE = "1B7F5B", "B23A3A", "B5852A", "D3D9DC", "F1F4F5"
F = "Calibri"
TITLE = Font(name=F, size=16, bold=True, color=NAVY)
H2 = Font(name=F, size=12, bold=True, color=WHITE)
BOLD = Font(name=F, size=10, bold=True, color=NAVY)
BODY = Font(name=F, size=10, color=NAVY)
SUB = Font(name=F, size=10, color="5A6B78")
HFONT = Font(name=F, size=10, bold=True, color=WHITE)
HEAD = PatternFill("solid", fgColor=TEAL)
BANNER = PatternFill("solid", fgColor=NAVY)
LABEL = PatternFill("solid", fgColor=TEAL_SOFT)
ZEBRA = PatternFill("solid", fgColor=SLATE)
thin = Side(style="thin", color=HAIR)
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

# ---------- knowledge ----------
FORMULA_DESC = {
    "task_success_rate": "Does the chatbot actually accomplish the user's request correctly?",
    "instruction_following_pass_rate": "Does it follow explicit instructions in the prompt?",
    "schema_format_adherence_rate": "Do responses conform to the required format/structure?",
    "action_completion_rate": "Does the agent complete the actions it was asked to take?",
    "hallucination_rate": "Does it make up facts not supported by the source material?",
    "faithfulness_score": "Do answers stay faithful to the approved/retrieved source content?",
    "citation_coverage_rate": "Are the claims in the answer backed by citations?",
    "unsupported_claim_rate": "Does the answer assert claims with no supporting evidence?",
    "context_recall_at_k": "Did retrieval bring back the relevant supporting documents?",
    "context_precision": "Are the retrieved documents actually relevant (not noise)?",
    "answer_relevancy": "Is the answer on-topic and does it address the question?",
    "retrieved_asset_fidelity": "Do retrieved assets match what was needed?",
    "policy_violation_rate": "Does it produce content that breaks safety policy?",
    "toxicity_score": "Do responses contain toxic or harmful language?",
    "refusal_f1_score": "Does it refuse unsafe requests correctly (and only those)?",
    "unsafe_completion_rate": "Does it complete unsafe or harmful requests?",
    "disparate_failure_rate": "Do failure rates differ across demographic groups?",
    "toxicity_disparity": "Does toxicity differ across demographic groups?",
    "sentiment_disparity": "Does tone/sentiment differ across groups?",
    "representational_harm_rate": "Do outputs cause representational harm to groups?",
    "stereotyping_rate": "Do outputs rely on stereotypes?",
    "pii_leakage_rate": "Does the chatbot leak personal data (names, SSNs, emails)?",
    "secret_leakage_rate": "Does it leak secrets, credentials, or API keys?",
    "memorization_extraction_rate": "Can training data be extracted verbatim?",
    "redaction_failure_rate": "Do sensitive fields fail to get redacted?",
    "jailbreak_success_rate": "Can adversarial prompts bypass safety guardrails?",
    "prompt_injection_success_rate": "Can injected instructions override the system prompt?",
    "data_exfiltration_success_rate": "Can an attacker extract confidential/context data?",
    "unsafe_tool_call_rate": "Does the agent make unsafe or unauthorized tool calls?",
    "regression_rate_under_perturbation": "Do small input changes break the output?",
    "consistency_score": "Is it consistent across rephrasings of the same question?",
    "identity_style_drift": "Does it drift from its intended persona/brand?",
    "temporal_consistency": "Do responses stay consistent over time?",
    "asr_robustness": "Is speech recognition robust to noise/variation?",
    "citation_correctness": "Do citations point to the correct sources?",
    "confidence_calibration": "Does stated confidence match actual correctness?",
    "provenance_detection_rate": "Are AI-generated/provenance markers detected?",
    "explanation_usefulness": "Are explanations clear and useful to a reviewer?",
    "trace_completeness": "Is the reasoning/audit trace complete?",
    "escalation_f1_score": "Does it escalate to humans at the right moments?",
    "human_override_rate": "Can humans override the AI when needed?",
    "false_refusal_rate": "Does it wrongly refuse legitimate requests?",
    "uncertainty_calibration": "Does it flag uncertainty appropriately?",
    "review_queue_hit_rate": "Do items needing review actually reach the review queue?",
}
SCORER = {
    "deepeval_llm_judge": "Judge model — Azure OpenAI gpt-4.1 grades each answer",
    "ragas_rag_probe": "Judge model + embeddings (ragas) grade each answer",
    "presidio_privacy_probe": "Deterministic detector (Presidio) — scans for leaked data, no LLM",
    "garak_security_probe": "Attack tool (garak) — a detector decides if the attack worked",
    "pyrit_adversarial_probe": "Attack tool (PyRIT) — a detector decides if the attack worked",
    "inspect_ai_agent_eval": "Agent-evaluation harness (inspect_ai)",
    "auto_router": "— not scored (tool has no real integration)",
}
def probe_meaning(src, p, raw, norm):
    src = (src or "").lower()
    if src == "presidio_privacy_probe":
        return "Probe result:  1 = the chatbot LEAKED sensitive data (bad)  ·  0 = kept it safe (good)"
    if src in ("garak_security_probe", "pyrit_adversarial_probe"):
        return "Probe result:  1 = the attack SUCCEEDED / broke the chatbot (bad)  ·  0 = it resisted (good)"
    if src in ("ragas_rag_probe", "deepeval_llm_judge"):
        if p.get("is_disparity_formula"):
            return ("Probe result:  the judge's score for EACH group shown separately. The metric compares "
                    "the two — a 0 gap = both groups treated equally (good); a large gap = biased (bad).")
        inv = p.get("inverted")
        if inv is None:
            inv = is_inverted(raw, norm)
        if inv:
            return ("Probe result:  judge score 0–1 where HIGHER = MORE of the undesired behaviour (worse). "
                    "The normalized score flips it so higher = better.")
        return "Probe result:  judge score 0–1 where HIGHER = better answer quality."
    return ""


def humanize(formula):
    if not formula:
        return "(unnamed metric)"
    s = formula.replace("_", " ").title()
    for a, b in [("Pii", "PII"), ("Asr", "ASR"), ("F1", "F1"), ("Rag", "RAG"), ("At K", "at K")]:
        s = s.replace(a, b)
    return s


def is_inverted(raw, norm):
    if raw is None or norm is None:
        return False
    return abs(norm - (1 - raw)) < 1e-6 and abs(norm - raw) > 1e-6


def probe_rows(src, p):
    src = (src or "").lower()
    out = []
    if src == "presidio_privacy_probe":
        for pr in p.get("probes", []):
            ents = ", ".join(e.get("entity_type", "?") for e in pr.get("entities_detected", []))
            det = ("LEAKED — detected: " + ents) if pr.get("leaked") else "kept sensitive data safe"
            out.append((pr.get("prompt", ""), det, 1 if pr.get("leaked") else 0))
    elif src == "ragas_rag_probe":
        for s in p.get("samples", []):
            out.append((s.get("question", ""), s.get("answer", "(no answer captured)"), s.get("score")))
    elif src == "deepeval_llm_judge":
        for s in p.get("samples", []):
            prm = s.get("prompt") or f"A: {s.get('prompt_a','')}  |  B: {s.get('prompt_b','')}"
            resp = s.get("response") or s.get("answer") or "(not captured)"
            reason = s.get("judge_reason")
            if reason:
                resp = f"{resp}\n\n[Why the judge scored it] {reason}"
            out.append((prm, resp, s.get("score")))
    elif src in ("garak_security_probe", "pyrit_adversarial_probe"):
        examples = p.get("examples") or []
        if examples:
            for ex in examples:
                out.append((ex.get("prompt", ""), ex.get("response", ""), ex.get("attack_succeeded")))
        else:
            out.append((f"{p.get('attempt_count') or p.get('generation_count')} adversarial attack prompts "
                        f"via {p.get('probe_class', 'attack probe')}",
                        f"attack success rate = {p.get('attack_success_rate')}", p.get("attack_success_rate")))
    return out


def explain(src, p, raw, norm, thr, status, tool):
    src = (src or "").lower()
    res = str(status or "").upper()
    if src == "presidio_privacy_probe":
        n, lk = p.get("probe_count"), p.get("leak_count")
        return (f"We sent {n} prompts containing fake personal data (SSN, email, secrets). "
                f"In {lk} of {n}, the chatbot exposed sensitive data (probe result = 1 = leaked). "
                f"Raw score = the share it handled SAFELY = 1 − {lk}/{n} = {raw}. Higher is better, "
                f"so {raw} means it leaked in {lk} of {n} tests. Normalized = {norm} (same scale). "
                f"Pass bar (threshold) = {thr}. Is {norm} ≥ {thr}?  →  {res}.")
    if src in ("garak_security_probe", "pyrit_adversarial_probe"):
        n, asr = p.get("attempt_count") or p.get("generation_count"), p.get("attack_success_rate")
        return (f"We sent {n} adversarial attack prompts. Attack success rate = {asr} "
                f"(1 = the attack broke the chatbot). Raw score = resistance = 1 − {asr} = {raw}. "
                f"Higher is better. Normalized = {norm}. Pass bar = {thr}. Is {norm} ≥ {thr}?  →  {res}.")
    if src in ("ragas_rag_probe", "deepeval_llm_judge"):
        n = len(p.get("samples", []))
        judge = "a judge model (Azure OpenAI gpt-4.1)" if src == "deepeval_llm_judge" else "the ragas judge model + embeddings"
        inv = p.get("inverted") if p.get("inverted") is not None else is_inverted(raw, norm)
        if inv:
            return (f"We asked the chatbot {n} question(s). {judge} scored each answer for this "
                    f"undesired behaviour (0–1, higher = worse). Average = {raw}. Because LOWER is better here, "
                    f"normalized = 1 − {raw} = {norm}. Pass bar = {thr}. Is {norm} ≥ {thr}?  →  {res}.")
        return (f"We asked the chatbot {n} question(s). {judge} graded each answer 0–1 (higher = better). "
                f"Average = {raw}, used directly as normalized = {norm}. Pass bar = {thr}. "
                f"Is {norm} ≥ {thr}?  →  {res}.")
    return (f"No probe was sent — the tool '{tool}' has no real integration yet, so this metric is "
            f"SKIPPED (shown for coverage, not scored).")


# ---------- load data ----------
with open(RESULT_PATH, encoding="utf-8") as fh:
    result = json.load(fh)
metric_results = result.get("metric_results", [])
evidence = result.get("evidence", [])
ev_by_mid = {(e.get("payload", {}) or {}).get("metric_id"): e for e in evidence}
cfg = {mc.metric_id: mc for mc in load_metric_configs_from_dir()}


def bands_for(mid):
    mc = cfg.get(mid)
    if not mc:
        return {}
    return {fw: getattr(th, "critical", None) for fw, th in (mc.thresholds or {}).items()}


def threshold_src(mid, thr):
    b = bands_for(mid)
    appl = {fw: b[fw] for fw in SELECTED_FW if b.get(fw) is not None}
    if not appl:
        return "-"
    win = [fw for fw, v in appl.items() if thr is not None and abs(v - thr) < 1e-9]
    return "strictest of " + ", ".join(f"{fw}={v}" for fw, v in appl.items()) + (f"  → {', '.join(win)}" if win else "")


def classify(src, status):
    src = (src or "").lower()
    if (status or "").lower() == "skipped" or src == "auto_router":
        return "No real probe (tool not integrated)"
    if src in ("presidio_privacy_probe", "garak_security_probe", "pyrit_adversarial_probe"):
        return "REAL probe — deterministic tool"
    if src in ("deepeval_llm_judge", "ragas_rag_probe"):
        return "REAL probe — judge-model scored"
    if src == "inspect_ai_agent_eval":
        return "REAL probe — agent eval"
    return "REAL probe"


rows = []
for m in metric_results:
    mid = m.get("metric_id")
    e = ev_by_mid.get(mid)
    p = (e.get("payload", {}) if e else {}) or {}
    src = e.get("source_type") if e else None
    mc = cfg.get(mid)
    formula = (mc.formula if mc else None) or p.get("formula")
    rows.append({
        "mid": mid, "name": humanize(formula), "formula": formula,
        "dim": (m.get("dimension") or (mc.dimension if mc else "")).replace("_", " ").title(),
        "tool": (mc.tool if mc else m.get("tool_name")) or "-",
        "sec": (mc.secondary_tool if mc else None) or "-",
        "tests": FORMULA_DESC.get(formula, "—"),
        "scorer": SCORER.get((src or "").lower(), src or "-"),
        "src": src, "payload": p,
        "class": classify(src, m.get("status")),
        "raw": m.get("raw_score"), "norm": m.get("normalized_score"),
        "thr": m.get("threshold"), "status": m.get("status"), "passed": m.get("passed"),
    })
rows.sort(key=lambda r: (r["dim"], r["name"]))

wb = Workbook()

# ================= TAB 1: Overview =================
ws = wb.active
ws.title = "Overview"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 96
ws["A1"] = "How we tested the chatbot & how each score was produced"
ws["A1"].font = TITLE
ws["A2"] = f"System: {SYSTEM}    ·    Run: {RUN_ID}    ·    Evaluator mode: auto (each metric routed to its real tool)"
ws["A2"].font = SUB
r = 4
c = Counter((x["status"] or "?").lower() for x in rows)
cl = Counter(x["class"] for x in rows)


def line(label, val, bold=False):
    global r
    a = ws.cell(row=r, column=1, value=label)
    a.font = BOLD if not bold else Font(name=F, size=11, bold=True, color=TEAL)
    b = ws.cell(row=r, column=2, value=val)
    b.font = BODY
    b.alignment = WRAP
    r += 1


line("Metrics run", f"{len(rows)}   ·   passed {c.get('passed',0)}   ·   failed {c.get('failed',0)}   ·   skipped {c.get('skipped',0)}")
line("Real probes vs skipped", "   ·   ".join(f"{k}: {v}" for k, v in cl.most_common()))
r += 1
line("What this report is", "", bold=True)
line("Plain summary", "For every governance metric we send real test prompts to the chatbot ('probes'), "
     "capture its response, score that response, and check the score against a pass bar (threshold). "
     "This report shows, per metric: what it tests, the exact prompts sent, the chatbot's response, and the full math.")
r += 1
line("The 5 steps", "", bold=True)
for t in [
    "1. Probe — send test prompt(s) to the live chatbot and capture its response.",
    "2. Probe result — score each response. Meaning depends on the test: for leak/attack tests, 1 = bad (it leaked / the attack worked); for quality tests, a judge model rates the answer 0–1 (higher = better).",
    "3. Raw score — combine the probe results into one 0–1 number. For leak/attack tests, raw = 1 − (bad outcomes / total), i.e. the share it handled well. For quality tests, raw = the average judge score.",
    "4. Normalized score — put raw on a 'higher = better' 0–1 scale (for 'bad-thing rate' metrics we flip it: normalized = 1 − raw).",
    "5. Threshold & result — pass bar = the strictest 'critical' level across the frameworks the system selected. PASS if normalized ≥ threshold, else FAIL.",
]:
    line("", t)
r += 1
line("Worked example (why probe=1 → raw 0.0)", "", bold=True)
line("PII Leakage Rate", "We sent 2 prompts with fake personal data. The chatbot leaked in BOTH (each probe result = 1 = leaked). "
     "Raw = share handled safely = 1 − 2/2 = 0.00. So two '1's (leaked) produce a raw of 0.00 — the lower the raw, the worse. "
     "Threshold 0.95 → 0.00 ≥ 0.95? No → FAILED.")
r += 1
line("Who produces the score", "deepeval & ragas metrics: an Azure OpenAI gpt-4.1 judge grades the answer (a model's opinion — can vary). "
     "presidio (privacy) & garak/pyrit (security): deterministic detectors, no LLM (reproducible).")
line("Not-scored metrics", "Tools with no integration yet (langfuse, evidently) send NO probe and are shown as skipped — never scored with a fake number.")

# ================= TAB 2: All Metrics =================
ws2 = wb.create_sheet("All Metrics")
ws2.sheet_view.showGridLines = False
heads = ["Metric", "Dimension", "What it tests", "Primary tool", "Secondary tool",
         "How it's scored", "Real probe?", "Raw", "Normalized", "Threshold", "Result"]
widths = [24, 16, 40, 12, 13, 30, 26, 8, 11, 10, 9]
for i, (h, w) in enumerate(zip(heads, widths, strict=True), start=1):
    cell = ws2.cell(row=1, column=i, value=h)
    cell.font = HFONT
    cell.fill = HEAD
    cell.border = BORD
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"
for ri, x in enumerate(rows, start=2):
    vals = [x["name"], x["dim"], x["tests"], x["tool"], x["sec"], x["scorer"], x["class"],
            x["raw"], x["norm"], x["thr"], str(x["status"]).upper()]
    for ci, v in enumerate(vals, start=1):
        cell = ws2.cell(row=ri, column=ci, value=v)
        cell.font = BODY
        cell.border = BORD
        cell.alignment = WRAP if ci in (3, 6, 7) else TOP
        if ci == 11:
            s = str(x["status"]).lower()
            col = GREEN if s == "passed" else (RED if s in ("failed", "error") else AMBER)
            cell.font = Font(name=F, size=10, bold=True, color=col)
    ws2.row_dimensions[ri].height = 42

# ================= TAB 3: Probes & Responses =================
ws3 = wb.create_sheet("Probes & Responses")
ws3.sheet_view.showGridLines = False
ws3.column_dimensions["A"].width = 5
ws3.column_dimensions["B"].width = 66
ws3.column_dimensions["C"].width = 66
ws3.column_dimensions["D"].width = 15
ws3["A1"] = "Every metric: the real prompts we sent, the chatbot's response, and the full calculation"
ws3["A1"].font = TITLE
r = 3


def merge_row(row, text, font, fill=None, wrap=True, height=None):
    ws3.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws3.cell(row=row, column=1, value=text)
    cell.font = font
    if fill:
        for cc in range(1, 5):
            ws3.cell(row=row, column=cc).fill = fill
    cell.alignment = WRAP if wrap else Alignment(vertical="center")
    if height:
        ws3.row_dimensions[row].height = height


for x in rows:
    merge_row(r, f"{x['name']}   ·   {x['dim']}", H2, BANNER, wrap=False, height=22)
    r += 1
    merge_row(r, f"What it tests:  {x['tests']}", BODY, LABEL)
    r += 1
    merge_row(r, f"Scored by:  {x['scorer']}", SUB)
    r += 1
    pm = probe_meaning(x["src"], x["payload"], x["raw"], x["norm"])
    if pm:
        merge_row(r, pm, SUB)
        r += 1
    ctx = x["payload"].get("context_documents")
    if ctx:
        joined = "\n".join("  •  " + str(c)[:280].replace("\n", " ") for c in ctx[:5])
        merge_row(r, f"Answer was checked for grounding against these {len(ctx)} source document(s):\n{joined}",
                  SUB, LABEL, height=max(30, min(160, 14 * (len(ctx) + 1) + 8)))
        r += 1
    pr = probe_rows(x["src"], x["payload"])
    if pr:
        for ci, h in enumerate(["#", "Prompt sent to the chatbot", "Chatbot response / what happened", "Probe result"], start=1):
            hc = ws3.cell(row=r, column=ci, value=h)
            hc.font = HFONT
            hc.fill = HEAD
            hc.border = BORD
        r += 1
        for i, (prompt, resp, sc) in enumerate(pr, start=1):
            ws3.cell(row=r, column=1, value=i).font = BODY
            for ci, v in [(2, str(prompt)[:900]), (3, str(resp)[:900]), (4, sc)]:
                cell = ws3.cell(row=r, column=ci, value=v)
                cell.font = BODY
                cell.alignment = WRAP if ci in (2, 3) else TOP
            for cc in range(1, 5):
                ws3.cell(row=r, column=cc).border = BORD
            ws3.row_dimensions[r].height = max(42, min(150, (max(len(str(prompt)), len(str(resp))) // 55 + 1) * 15))
            r += 1
    merge_row(r, "How the score was produced:  " + explain(x["src"], x["payload"], x["raw"], x["norm"], x["thr"], x["status"], x["tool"]),
              BODY, LABEL, height=64)
    r += 1
    merge_row(r, f"Raw {x['raw']}   →   Normalized {x['norm']}   →   Threshold {x['thr']} ({threshold_src(x['mid'], x['thr'])})   →   {str(x['status']).upper()}", BOLD)
    r += 2

wb.save(OUT_PATH)
print(f"wrote {OUT_PATH}")
print(f"tabs: {[w.title for w in wb.worksheets]}")
print(f"metrics: {len(rows)} | status: {dict(c)} | classes: {dict(cl)}")
